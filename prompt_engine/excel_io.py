import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment

def load_input_excel(path):
    df = pd.read_excel(path)
    if "Textfeld" not in df.columns:
        raise ValueError("Die Excel-Datei muss eine Spalte namens 'Textfeld' enthalten.")
    return df

def save_results_excel(results, path):
    df = pd.DataFrame(results)
    df.to_excel(path, index=False)

def format_excel_columns(path):
    wb = load_workbook(path)
    ws = wb.active

    for col in ws.columns:
        max_length = 0
        col_letter = get_column_letter(col[0].column)

        for cell in col:
            cell.alignment = Alignment(wrap_text=True)
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))

        ws.column_dimensions[col_letter].width = min(max_length * 0.9, 100)

    wb.save(path)
